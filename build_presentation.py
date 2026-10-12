"""
Generate Widescreen Academic Presentation (.pptx)
"Inside an AI's Mind: When Knowledge Becomes a Graph"
CIA III-B Mini Project — Graph Theory in My World (Domain: AI)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette
    BG_DARK = RGBColor(11, 15, 25)         # Deep Obsidian #0B0F19
    CARD_BG = RGBColor(22, 29, 47)         # Dark Slate Blue #161D2F
    CARD_BORDER = RGBColor(45, 55, 80)     # Subtle border #2D3750
    TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8
    TEXT_BODY = RGBColor(203, 213, 225)    # #CBD5E1

    ACCENT_GOLD = RGBColor(245, 158, 11)   # Core AI / Gold #F59E0B
    ACCENT_GREEN = RGBColor(16, 185, 129)  # Speaker A / Emerald #10B981
    ACCENT_ROSE = RGBColor(244, 63, 94)    # Speaker B / Rose #F43F5E
    ACCENT_CYAN = RGBColor(6, 182, 212)    # Learning / Cyan #06B6D4
    ACCENT_VIOLET = RGBColor(139, 92, 246) # Special / Violet #8B5CF6

    FONT_FAMILY = "Segoe UI"

    slides_data = [
        # Slide 1: Title
        {
            "slide_num": 1,
            "category": "CIA III-B MINI PROJECT • GRAPH THEORY IN MY WORLD",
            "title": "Inside an AI's Mind: When Knowledge Becomes a Graph",
            "speaker": "TEAM OF 2 • 10-MINUTE ORAL DEFENSE",
            "speaker_color": ACCENT_GOLD,
            "time_badge": "0:30 (Elapsed 0:30)",
            "bullets": [
                ("Creative Artifact Domain:", "Artificial Intelligence modelled as an explicit knowledge graph."),
                ("Graph Theory is the Hero:", "We show how discrete mathematics — vertices, edges, degrees, paths, cycles, and connectivity — reveals structure inside an AI's mind."),
                ("Discrete Graph Model:", "A simple, undirected graph G = (V, E) with |V| = 28 vertices and |E| = 39 edges."),
                ("Strict Academic Discipline:", "No actual ML models simulated, no future commercialization claims, pure computed graph theory.")
            ],
            "takeaway": "“Graph Theory is the hero; Artificial Intelligence is only our creative world.”",
            "image": "screenshots/3d_universe.png",
            "image_caption": "Interactive 3D Three.js Knowledge Graph Universe with OrbitControls & CSS2D Billboards"
        },
        # Slide 2: Why AI is Our Personal World
        {
            "slide_num": 2,
            "category": "PROBLEM FORMULATION & SCOPE",
            "title": "Why AI is Our Personal 'World' & Mathematical Scope",
            "speaker": "SPEAKER A",
            "speaker_color": ACCENT_GREEN,
            "time_badge": "0:45 (Elapsed 1:15)",
            "bullets": [
                ("Demystifying the 'Black Box':", "AI is commonly treated as opaque matrices of numbers. Graph theory offers structural interpretability and conceptual clarity."),
                ("Relational Graph Metaphor:", "Knowledge concepts become vertices; informational dependencies become undirected edges."),
                ("Mathematical Constraints:", "Simple graph (no self-loops, no parallel/multi-edges). Undirected edges represent bidirectional conceptual access."),
                ("Discrete Rigor:", "Every theorem, degree count, path length, and component number is strictly computed from verified ground truth.")
            ],
            "takeaway": "Graph Theory transforms an opaque black box into an explainable, structured relational universe.",
            "image": "screenshots/2d_dashboard.png",
            "image_caption": "2D Analytics Dashboard with live Discrete Math KPI Strip (|V|=28, |E|=39, Σdeg=78, c=2)"
        },
        # Slide 3: Transforming Knowledge into a Graph
        {
            "slide_num": 3,
            "category": "GRAPH TOPOLOGY & FORMAL SPECIFICATION",
            "title": "Transforming Knowledge into a Graph: G = (V, E)",
            "speaker": "SPEAKER A",
            "speaker_color": ACCENT_GREEN,
            "time_badge": "1:00 (Elapsed 2:15)",
            "bullets": [
                ("Order & Size of Graph G:", "|V| = 28 vertices, |E| = 39 edges across 6 distinct semantic domains."),
                ("Six Domain Clusters:", "Core AI (1), Learning & Data (11), Vision (4), NLP (6), Reasoning (3), Specialized Chess (3)."),
                ("Two Connected Components:", "Main Knowledge Network K1 (25 vertices, 37 edges) and Specialized Cluster K2 (3 vertices, 2 edges)."),
                ("Specialized Disconnection:", "The chess rules cluster (chessrules ━ gametree ━ strategy) has zero paths to the main network, proving c = 2.")
            ],
            "takeaway": "G consists of 2 components: a cohesive main mind (K1) and an isolated specialist cluster (K2).",
            "image": "screenshots/01_normal.png",
            "image_caption": "Topological map of G = (V, E) showing the 25-node main mind and 3-node chess cluster"
        },
        # Slide 4: Vertices and Edges in Action
        {
            "slide_num": 4,
            "category": "RELATIONAL ONTOLOGY & BIDIRECTIONALITY",
            "title": "Vertices, Edges, and Bidirectional Relational Semantics",
            "speaker": "SPEAKER A",
            "speaker_color": ACCENT_GREEN,
            "time_badge": "0:45 (Elapsed 3:00)",
            "bullets": [
                ("Vertex Definition:", "Each vertex v in V encapsulates a discrete knowledge concept with an exact <= 10 word meaning."),
                ("Edge Definition:", "An edge e = {u, v} represents an undirected mutual relationship. Conceptual queries operate both ways."),
                ("Case Study (E13: Training ━ Model):", "Training shapes and parameters the Model; the Model architecture conversely defines training requirements."),
                ("Incidence & Adjacency:", "The vertex 'training' has degree 4, incident to edges connecting it to 'ml', 'nn', 'model', and 'data'.")
            ],
            "takeaway": "Knowledge is inherently bidirectional: learning adapts models, and models guide future learning.",
            "image": "screenshots/09_theory_drawer.png",
            "image_caption": "Interactive Graph Theory Concept Cards Drawer explaining Vertex and Edge foundations"
        },
        # Slide 5: Vertex Degree & The Handshaking Lemma
        {
            "slide_num": 5,
            "category": "FUNDAMENTAL THEOREM & DEGREE DISTRIBUTION",
            "title": "Vertex Degree & The Handshaking Lemma",
            "speaker": "SPEAKER A",
            "speaker_color": ACCENT_GREEN,
            "time_badge": "0:45 (Elapsed 3:45)",
            "bullets": [
                ("Degree Definition:", "deg(v) equals the number of incident edges. deg(v) = |{u in V : {u, v} in E}|."),
                ("Euler's Handshaking Lemma:", "Sum of all degrees = 2|E|. For G: 2 * 39 = 78. (K1 sum = 74 = 2*37; K2 sum = 4 = 2*2)."),
                ("Surprising Architectural Finding:", "'Data' is the unique maximal degree hub (deg = 7), NOT 'AI' (deg = 5)."),
                ("Degree Sequence of K1:", "[7, 5, 5, 5, 4, 4, 4, 3, 3, 3, 3, 3, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 1]. All 28 degrees verified.")
            ],
            "takeaway": "Discrete Math proves the AI's mind is fundamentally data-driven, not executive-driven.",
            "image": "screenshots/02_degree.png",
            "image_caption": "Ranked degree distribution verifying Data as maximal hub (deg 7) and Handshaking sum = 78"
        },
        # Slide 6: Paths, Trails, and Reasoning Chains
        {
            "slide_num": 6,
            "category": "CONNECTIVITY & REASONING PIPELINES",
            "title": "Paths, Trails, and Reasoning Chains",
            "speaker": "SPEAKER B",
            "speaker_color": ACCENT_ROSE,
            "time_badge": "1:00 (Elapsed 4:45)",
            "bullets": [
                ("Walk vs Trail vs Path:", "A Walk allows repeats; a Trail forbids repeated edges; a Path forbids repeated vertices."),
                ("Path P1 (Inference Pipeline):", "Data -> Training -> Model -> Prediction (length 3 edges, 4 distinct vertices)."),
                ("Path P2 (NLP Semantic Grounding):", "Words -> Meaning -> Context -> Knowledge (length 3 edges, 4 distinct vertices)."),
                ("Shortest Path (P3):", "Dataset -> Data -> AI -> Knowledge -> Decision (length 4 edges). Dataset -> Strategy has NO path (null).")
            ],
            "takeaway": "Logical inference in an AI corresponds strictly to simple paths between conceptual vertices.",
            "image": "screenshots/03_path_p1.png",
            "image_caption": "Path P1 highlighted in brilliant emerald across the 3D neural topography"
        },
        # Slide 7: Cycles & Feedback Loops in Learning
        {
            "slide_num": 7,
            "category": "CYCLIC TOPOLOGY & LEARNING LOOPS",
            "title": "Cycles & Feedback Loops in Learning",
            "speaker": "SPEAKER B",
            "speaker_color": ACCENT_ROSE,
            "time_badge": "0:45 (Elapsed 5:30)",
            "bullets": [
                ("Cycle Definition:", "A closed walk of length >= 3 where start = end, with no repeated intermediate vertices or edges."),
                ("Primary Learning Cycle C1 (Length 5):", "ML -> Model -> Prediction -> Feedback -> Data -> ML. Models supervised retraining."),
                ("Structural Triangles (K3 Cliques):", "C2: ML -> Training -> Model -> ML. C3: Language -> Words -> Text -> Language."),
                ("Acyclic Trees vs Cyclic Minds:", "If knowledge were an acyclic tree, an AI could never learn from feedback. Cycles are vital.")
            ],
            "takeaway": "Without cycles, an AI cannot learn from mistakes: feedback loops require cyclic graph topology.",
            "image": "screenshots/04_cycle_c1.png",
            "image_caption": "Cycle C1 feedback loop illuminated in vivid cyan, tracing model adaptation"
        },
        # Slide 8: Subgraphs vs. Components: Domain Clusters
        {
            "slide_num": 8,
            "category": "SUBGRAPH THEORY & MODULAR ARCHITECTURE",
            "title": "Subgraphs vs. Connected Components: Domain Clusters",
            "speaker": "SPEAKER B",
            "speaker_color": ACCENT_ROSE,
            "time_badge": "0:45 (Elapsed 6:15)",
            "bullets": [
                ("Subgraph Definition:", "H = (V_H, E_H) where V_H subset of V and E_H subset of E. An induced subgraph preserves all mutual edges."),
                ("Common Misconception Debunked:", "Domain modules (Vision, NLP, Learning) are subgraphs, NOT components, because boundary edges tie them to G."),
                ("Featured Subgraphs:", "S1 Learning (11 nodes, 16 edges, 5 boundary edges); S2 Vision (4 nodes, 3 edges, 3 boundary edges)."),
                ("Vision Boundary Edges:", "E2 (AI ━ CV), E25 (Recognition ━ Knowledge), and E26 (Image ━ Data) integrate seeing with thinking.")
            ],
            "takeaway": "Specialised domains are modular subgraphs anchored to the main core by explicit boundary edges.",
            "image": "screenshots/05_subgraph_s2.png",
            "image_caption": "Computer Vision Subgraph S2 highlighted in rose with gold boundary edge connectors"
        },
        # Slide 9: Connectedness, Cut Vertices & Bridges
        {
            "slide_num": 9,
            "category": "RESILIENCE, BOTTLENECKS & VULNERABILITIES",
            "title": "Connectedness, Cut Vertices & Bridges: System Vulnerabilities",
            "speaker": "SPEAKER B",
            "speaker_color": ACCENT_ROSE,
            "time_badge": "1:00 (Elapsed 7:15)",
            "bullets": [
                ("Connected Component:", "A maximal connected subgraph. Whole graph G has exactly 2 components (c = 2)."),
                ("Cut Vertex Discovery:", "'Data' is a Cut Vertex: removing 'Data' increases components from 2 to 3, stranding 'Dataset' as an isolated node!"),
                ("Is 'AI' a Cut Vertex?:", "NO! Removing 'AI' leaves K1 connected (c remains 2) due to alternate redundant circuits."),
                ("Bridge Edge E18:", "Edge E18 (Data ━ Dataset) is a Bridge: removing E18 instantly increases components from 2 to 3.")
            ],
            "takeaway": "Data is the single point of catastrophic failure; the AI hub itself has resilient redundant circuits.",
            "image": "screenshots/06_component_remove_data.png",
            "image_caption": "Component fragmentation demonstration: removing Data increases component count from 2 to 3"
        },
        # Slide 10: Critical Thinking, Limitations & Analysis
        {
            "slide_num": 10,
            "category": "EVALUATION, REFLECTION & CRITICAL THINKING",
            "title": "Critical Thinking, Limitations & Analysis",
            "speaker": "SPEAKER A & B",
            "speaker_color": ACCENT_GOLD,
            "time_badge": "0:45 (Elapsed 8:00)",
            "bullets": [
                ("What Graph Theory Exposes:", "Uncovers architectural vulnerabilities (cut vertices) and reasoning pathways that matrix weights conceal."),
                ("Model Limitations Acknowledged:", "Static topology (real neural embeddings adapt dynamically); unweighted edges (treats all bonds equally)."),
                ("Bidirectional Simplification:", "Undirected edges represent associative access, though causality in deep networks is often directional."),
                ("Academic Conclusion:", "Discrete graph theory is a legitimate, rigorous, and intuitive framework for explainable AI.")
            ],
            "takeaway": "Discrete graph theory provides an elegant formal bridge between human reasoning and artificial systems.",
            "image": "screenshots/10_analysis_drawer.png",
            "image_caption": "Observations & Critical Thinking Drawer evaluating model strengths and limitations"
        },
        # Slide 11: Summary, Interactive Artifact & Viva Q&A
        {
            "slide_num": 11,
            "category": "CONCLUSION & VIVA VOCE EXAMINATION",
            "title": "Summary, Interactive System & Viva Q&A",
            "speaker": "SPEAKER A & B",
            "speaker_color": ACCENT_GOLD,
            "time_badge": "1:00 (Total 09:00 + 1:00 Buffer)",
            "bullets": [
                ("Summary of Verified Metrics:", "|V|=28, |E|=39, Components c=2, Handshaking Σdeg=78, Max Hub: Data (7), Cut Vertex: Data, Bridge: E18."),
                ("Live Interactive Suite:", "3D Three.js Universe + 2D Vector Analytics Dashboard + AI Mind Lab (4 Missions & Glitch Simulation)."),
                ("Examiner Rapid Demo Ready:", "Examiner mode preloaded with one-click viva demonstrations for instant proof of all theorems."),
                ("Thank You!", "We welcome questions from the examiners and faculty on any aspect of the graph topology!")
            ],
            "takeaway": "All 31 discrete mathematics assertions verified 100%. Live interactive demo ready for examination.",
            "image": "screenshots/presentation_deck.png",
            "image_caption": "Slide-driven interactive presentation HUD with synchronized live graph animations"
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # 1. Background Fill: Solid Deep Obsidian
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_DARK
        bg_shape.line.fill.background()

        # 2. Header Box (Left: Category + Title; Right: Speaker + Time Badge)
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = data["category"]
        p_cat.font.name = FONT_FAMILY
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = data["speaker_color"]

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(8.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Speaker & Time Badges (Top Right)
        badge_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.6), Inches(0.5), Inches(2.9), Inches(0.65)
        )
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = CARD_BG
        badge_box.line.color.rgb = data["speaker_color"]
        badge_box.line.width = Pt(1.5)
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True
        p_spk = tf_badge.paragraphs[0]
        p_spk.text = data["speaker"]
        p_spk.alignment = PP_ALIGN.CENTER
        p_spk.font.name = FONT_FAMILY
        p_spk.font.size = Pt(11)
        p_spk.font.bold = True
        p_spk.font.color.rgb = data["speaker_color"]

        p_time = tf_badge.add_paragraph()
        p_time.text = data["time_badge"]
        p_time.alignment = PP_ALIGN.CENTER
        p_time.font.name = FONT_FAMILY
        p_time.font.size = Pt(9.5)
        p_time.font.color.rgb = TEXT_MUTED

        # 3. Left Content Card (Text & Mathematical Explanations)
        left_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(6.4), Inches(5.1)
        )
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = CARD_BG
        left_card.line.color.rgb = CARD_BORDER
        left_card.line.width = Pt(1)

        # Content Text Box inside Left Card
        content_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(6.0), Inches(3.7))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for i, (b_lead, b_desc) in enumerate(data["bullets"]):
            p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
            p.space_after = Pt(10)
            
            # Bullet point bullet lead (Colored / Bold)
            run_lead = p.add_run()
            run_lead.text = f"•  {b_lead} "
            run_lead.font.name = FONT_FAMILY
            run_lead.font.size = Pt(12)
            run_lead.font.bold = True
            run_lead.font.color.rgb = data["speaker_color"]

            # Bullet description
            run_desc = p.add_run()
            run_desc.text = b_desc
            run_desc.font.name = FONT_FAMILY
            run_desc.font.size = Pt(11.5)
            run_desc.font.color.rgb = TEXT_BODY

        # Takeaway Callout Box at bottom of Left Card
        takeaway_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.65), Inches(6.0), Inches(0.85)
        )
        takeaway_card.fill.solid()
        takeaway_card.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
        takeaway_card.line.color.rgb = data["speaker_color"]
        takeaway_card.line.width = Pt(1)
        tf_takeaway = takeaway_card.text_frame
        tf_takeaway.word_wrap = True
        p_ta_tag = tf_takeaway.paragraphs[0]
        p_ta_tag.text = "KEY TAKEAWAY"
        p_ta_tag.font.name = FONT_FAMILY
        p_ta_tag.font.size = Pt(8.5)
        p_ta_tag.font.bold = True
        p_ta_tag.font.color.rgb = data["speaker_color"]

        p_ta = tf_takeaway.add_paragraph()
        p_ta.text = data["takeaway"]
        p_ta.font.name = FONT_FAMILY
        p_ta.font.size = Pt(10.5)
        p_ta.font.bold = True
        p_ta.font.color.rgb = TEXT_WHITE

        # 4. Right Visual Card (Embedded Screenshot)
        right_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(1.6), Inches(5.1), Inches(5.1)
        )
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = CARD_BG
        right_card.line.color.rgb = CARD_BORDER
        right_card.line.width = Pt(1)

        # Embed Screenshot Image if available
        img_path = data["image"]
        if os.path.exists(img_path):
            try:
                # Add picture inside the right card with margin
                slide.shapes.add_picture(
                    img_path, Inches(7.55), Inches(1.75), width=Inches(4.8), height=Inches(3.8)
                )
            except Exception as e:
                print(f"Error embedding picture {img_path}: {e}")

        # Visual Caption Text
        caption_box = slide.shapes.add_textbox(Inches(7.55), Inches(5.7), Inches(4.8), Inches(0.8))
        tf_cap = caption_box.text_frame
        tf_cap.word_wrap = True
        p_cap = tf_cap.paragraphs[0]
        p_cap.text = f"Figure {data['slide_num']}: {data['image_caption']}"
        p_cap.alignment = PP_ALIGN.CENTER
        p_cap.font.name = FONT_FAMILY
        p_cap.font.size = Pt(9.5)
        p_cap.font.italic = True
        p_cap.font.color.rgb = TEXT_MUTED

        # 5. Footer Line (Mandatory Academic Disclaimer & Slide Counter)
        foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4))
        tf_foot = foot_box.text_frame
        tf_foot.word_wrap = True
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = (
            "A conceptual graph model used to demonstrate how Graph Theory can represent relationships "
            "between knowledge concepts. It is not how a real AI stores knowledge internally.  |  "
            f"Slide {data['slide_num']} of {len(slides_data)}"
        )
        p_foot.font.name = FONT_FAMILY
        p_foot.font.size = Pt(8.5)
        p_foot.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    output_filename = "Inside_An_AIs_Mind_Presentation.pptx"
    prs.save(output_filename)
    print(f"Successfully generated PowerPoint presentation: {output_filename} ({len(slides_data)} slides)")

if __name__ == "__main__":
    create_deck()
